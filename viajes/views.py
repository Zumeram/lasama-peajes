from django.shortcuts import render,redirect,get_object_or_404
from django.conf import settings
from django.core.mail import EmailMessage
from django.utils import timezone
from .models import Viaje,Peaje
import openpyxl,io
from openpyxl.styles import Font
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from datetime import datetime
def lista_viajes(r):
    return render(r,'viajes/lista.html',{'viajes':Viaje.objects.all().order_by('-fecha_inicio')})
def iniciar_viaje(r):
    if r.method=='POST':
        o=r.POST.get('origen','Rengo'); d=r.POST.get('destino','Chiloé')
        cod=f"LASAMA-{o[:4].upper()}-{d[:4].upper()}-{datetime.now().strftime('%Y%m%d')}-{Viaje.objects.count()+1:03d}"
        v=Viaje.objects.create(codigo=cod,origen=o,destino=d)
        return redirect('detalle_viaje',codigo=cod)
    return render(r,'viajes/iniciar.html')
def detalle_viaje(r,codigo):
    return render(r,'viajes/detalle.html',{'viaje':get_object_or_404(Viaje,codigo=codigo)})
def agregar_peaje(r,codigo):
    v=get_object_or_404(Viaje,codigo=codigo)
    if r.method=='POST':
        Peaje.objects.create(viaje=v,tipo=r.POST.get('tipo','TRONCAL'),nombre=r.POST.get('nombre','Peaje'),monto=int(r.POST.get('monto',0)),imagen=r.FILES.get('imagen'))
        return redirect('detalle_viaje',codigo=codigo)
    return render(r,'viajes/agregar.html',{'viaje':v})
def finalizar_viaje(r,codigo):
    v=get_object_or_404(Viaje,codigo=codigo); v.finalizado=True; v.save()
    wb=openpyxl.Workbook(); ws=wb.active; ws['A1']=f"LASAMA {codigo}"; tot=0; row=3
    for p in v.peajes.all():
        ws[f'A{row}']=p.tipo; ws[f'B{row}']=p.nombre; ws[f'C{row}']=p.monto; tot+=p.monto; row+=1
    ws[f'B{row+1}']='TOTAL'; ws[f'C{row+1}']=tot
    eb=io.BytesIO(); wb.save(eb); eb.seek(0)
    pb=io.BytesIO(); c=canvas.Canvas(pb,pagesize=letter); c.drawString(50,750,f"{codigo} ${tot}"); c.showPage(); c.save(); pb.seek(0)
    try:
        em=EmailMessage(f"[LASAMA] {codigo} ${tot}",f"Viaje {codigo} total ${tot}",settings.DEFAULT_FROM_EMAIL,[settings.FINANZAS_EMAIL])
        em.attach(f"{codigo}.xlsx",eb.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        em.attach(f"{codigo}.pdf",pb.getvalue(),'application/pdf'); em.send()
        msg=f"Enviado a {settings.FINANZAS_EMAIL}"
    except Exception as e: msg=str(e)
    return render(r,'viajes/finalizado.html',{'viaje':v,'total':tot,'mensaje':msg})
